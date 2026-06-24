"""
Tinte solar ETL — raw/Tinte/Tinte.xlsx -> data/tinte/*.parquet (+ meta.json).

Tinte is a ~10 MW AC solar PV plant (Netherlands, "Ampyr Tinte"). The workbook
is a one-month (April 2026) SCADA/technical export: weather (irradiance + module
temperature), plant + per-inverter AC power, PPC setpoint (curtailment), string
DC currents, and the grid main-meter. This module materialises clean series and
daily aggregates the Tinte Solar KPI dashboard reads.

Run when the Tinte workbook changes:
    python -m src.data_cleaning.process_tinte
"""
from __future__ import annotations

import datetime as _dt
import json
import os
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "raw" / "Tinte" / "Tinte.xlsx"
OUT = ROOT / "data" / "tinte"

CAP_AC_MW = 9.8           # nameplate AC (monthly report; PPC setpoint ceiling 10,000 kW)
CAP_DC_KWP = 12600.0      # DC nameplate (kWp) — Apr-26 technical report (12.6 MWp)

# April-2026 figures from AM_Monthly_Technical_Performance_Report-Apr26.pdf
# (Tinte section, pp.10-12). Budget = the report's "Budgeted" baseline (not P50).
REPORT_APR26 = {
    "month": "April 2026",
    "actual_gen_mwh": 1343,
    "budget_gen_mwh": 1410,
    "budget_max_theoretical_mwh": 1801,
    "pct_of_budget": 95.2,
    "actual_radiation_kwh_m2": 161,
    "budget_radiation_kwh_m2": 143,
    # Energy waterfall that reconciles theoretical-max -> actual (MWh), report p.11
    "energy_waterfall": [
        ("Theoretical max", 2029),
        ("Temperature", 73),
        ("Array loss", -473),
        ("Inverter loss", -18),
        ("AC loss", -30),
        ("Plant breakdown", -90),
        ("Grid + curtailment", -148),
    ],  # cumulative end = 1343 (Actual)
    "strings": {"total": 935, "normal": 774, "underperforming": 143, "faulty": 17, "severe": 1},
    "worst_strings": [
        ("1.18.DC4", -21.5), ("1.17.DC12", -13.7), ("1.15.DC14", -12.9),
        ("1.16.DC18", -12.4), ("1.15.DC13", -11.5),
    ],
    "note": "Grid and curtailment losses are combined and 'tentative' per the report; "
            "PR%, availability%, specific yield and degradation are not reported.",
}


def _isdt(x):
    return isinstance(x, _dt.datetime)


def _weather(wb) -> pd.DataFrame:
    ws = wb["WMS"]
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        t, irr, temp = r[0], r[1], r[2]
        if _isdt(t):
            rows.append((t, irr if isinstance(irr, (int, float)) else np.nan,
                         temp if isinstance(temp, (int, float)) else np.nan))
    df = pd.DataFrame(rows, columns=["time", "irradiance_wm2", "temp_module_c"]).set_index("time")
    return df.resample("15min").mean()


def _power(wb) -> pd.DataFrame:
    """Plant total AC power (kW) at 15-min from the 'sum' column."""
    ws = wb["Active Power 15 min"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        t, s = r[0], r[2]
        if _isdt(t):
            rows.append((t, s if isinstance(s, (int, float)) else np.nan))
    return pd.DataFrame(rows, columns=["time", "power_kw"]).set_index("time")


def _setpoint(wb) -> pd.DataFrame:
    """PPC setpoint % (col 4, 'aa'): 100 = unconstrained, <100 = curtailment."""
    ws = wb["Setpoint"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        t, pct = r[0], r[3]
        if _isdt(t):
            rows.append((t, pct if isinstance(pct, (int, float)) else np.nan))
    df = pd.DataFrame(rows, columns=["time", "setpoint_pct"]).set_index("time")
    return df.resample("15min").min()


def _meter(wb) -> pd.DataFrame:
    """Main-meter metered energy per 15-min interval (kWh) from 'Verbruik totaal'."""
    ws = wb["Meter data"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        t_to, vtot = r[3], r[6]
        if _isdt(t_to):
            rows.append((t_to, vtot if isinstance(vtot, (int, float)) else np.nan))
    df = pd.DataFrame(rows, columns=["time", "export_kwh"]).set_index("time").sort_index()
    return df


def _inverter_daily(wb) -> pd.DataFrame:
    """Per-inverter daily average AC power (kW) — relative comparison only."""
    ws = wb["Daily AC Power"]
    header = None
    data = []
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True)):
        if i == 0:
            header = [str(c) if c is not None else "" for c in r]
            continue
        if not _isdt(r[0]):
            continue
        rec = {"date": r[0]}
        for j in range(1, len(header)):
            name = header[j].replace("Ampyr Tinte.", "").replace(" - Power AC", "").strip()
            if name:
                rec[name] = r[j] if isinstance(r[j], (int, float)) else np.nan
        data.append(rec)
    df = pd.DataFrame(data)
    if "date" in df:
        df["date"] = pd.to_datetime(df["date"]).dt.date
    return df


def run(verbose: bool = True) -> dict:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    weather = _weather(wb)
    power = _power(wb)
    setp = _setpoint(wb)
    meter = _meter(wb)
    inv_daily = _inverter_daily(wb)

    # 15-min unified ops frame
    ops = weather.join(power, how="outer").join(setp, how="outer")
    ops.index.name = "time"
    ops.to_parquet(OUT / "ops_15min.parquet")
    meter.to_parquet(OUT / "meter_15min.parquet")
    inv_daily.to_parquet(OUT / "inverter_daily.parquet")

    # ---- Daily aggregates ----
    day = pd.DataFrame(index=pd.Index(sorted({t.date() for t in weather.index}), name="date"))

    # metered energy per day (MWh) — grid ground truth
    md = meter.copy()
    md["d"] = [t.date() for t in md.index]
    energy_kwh = md.groupby("d")["export_kwh"].sum()
    day["energy_mwh"] = (energy_kwh / 1000.0).reindex(day.index)

    # insolation kWh/m2 per day = sum(W/m2 * 0.25h) / 1000
    w = weather.copy()
    w["d"] = [t.date() for t in w.index]
    insol = w.groupby("d")["irradiance_wm2"].apply(lambda s: s.fillna(0).mul(0.25).sum() / 1000.0)
    day["insolation_kwh_m2"] = insol.reindex(day.index)
    day["avg_module_temp_c"] = w.groupby("d")["temp_module_c"].mean().reindex(day.index)

    # power-derived peak + curtailment flag
    p = power.copy()
    p["d"] = [t.date() for t in p.index]
    day["peak_mw"] = (p.groupby("d")["power_kw"].max() / 1000.0).reindex(day.index)
    sp = setp.copy()
    sp["d"] = [t.date() for t in sp.index]
    day["min_setpoint_pct"] = sp.groupby("d")["setpoint_pct"].min().reindex(day.index)
    day["curtail_intervals"] = sp[sp["setpoint_pct"] < 99.5].groupby("d")["setpoint_pct"].count().reindex(day.index).fillna(0)

    # yields + PR
    day["specific_yield_kwh_kwp"] = day["energy_mwh"] * 1000.0 / CAP_DC_KWP
    day["pr"] = (day["specific_yield_kwh_kwp"] / day["insolation_kwh_m2"]).replace([np.inf, -np.inf], np.nan)

    # uptime proxy: during daylight (irr>50 W/m2), fraction of 15-min intervals
    # where the plant was producing (power > 1% of AC). True availability needs
    # inverter status flags (not in this feed) — this is a coarse plant-uptime.
    dl = ops[ops["irradiance_wm2"] > 50].copy()
    if len(dl):
        dl["d"] = [t.date() for t in dl.index]
        dl["up"] = (dl["power_kw"] > CAP_AC_MW * 1000 * 0.01).astype(float)
        day["uptime_proxy_pct"] = (dl.groupby("d")["up"].mean() * 100).reindex(day.index)

    day.to_parquet(OUT / "daily.parquet")

    inv_cols = [c for c in inv_daily.columns if c != "date"]
    meta = {
        "asset": "Ampyr Tinte",
        "type": "Solar PV",
        "country": "Netherlands",
        "capacity_ac_mw": CAP_AC_MW,
        "capacity_dc_kwp": CAP_DC_KWP,
        "n_inverters": len(inv_cols),
        "period_start": str(min(day.index)),
        "period_end": str(max(day.index)),
        "n_days": int(day["energy_mwh"].notna().sum()),
        "total_energy_mwh": round(float(day["energy_mwh"].sum()), 1),
        "total_insolation_kwh_m2": round(float(day["insolation_kwh_m2"].sum()), 1),
        "specific_yield_month_kwh_kwp": round(float(day["energy_mwh"].sum() * 1000 / CAP_DC_KWP), 1),
        "avg_pr": round(float(day["pr"].mean()), 3),
        "peak_mw": round(float(day["peak_mw"].max()), 2),
        "avg_module_temp_c": round(float(day["avg_module_temp_c"].mean()), 1),
        "curtailment_intervals": int(day["curtail_intervals"].sum()),
        "report": REPORT_APR26,
    }
    with open(OUT / "meta.json", "w", encoding="utf-8") as fh:
        json.dump(meta, fh, indent=2)

    if verbose:
        print("Tinte ETL ->", OUT)
        for k, v in meta.items():
            print(f"  {k:28} {v}")
    return meta


# ---- readers ----
def _read(name):
    p = OUT / name
    return pd.read_parquet(p) if p.exists() else pd.DataFrame()


def read_daily():
    return _read("daily.parquet")


def read_ops_15min():
    return _read("ops_15min.parquet")


def read_meter_15min():
    return _read("meter_15min.parquet")


def read_inverter_daily():
    return _read("inverter_daily.parquet")


def read_meta():
    p = OUT / "meta.json"
    if p.exists():
        with open(p, encoding="utf-8") as fh:
            return json.load(fh)
    return {}


if __name__ == "__main__":
    run()
